# -*- coding: utf-8 -*-
"""부서 업무일지 v3 → v4 업그레이드 · 공통 설정"""
import datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

SRC = "부서_업무일지_v3.xlsx"
OUT = "부서_업무일지_v4.xlsx"
OUT2 = "부서_업무일지_v4.xlsx"

F = "Arial"
NAVY = "1F3864"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
SEC_FILL = PatternFill("solid", fgColor="2E75B6")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
TOT_FILL = PatternFill("solid", fgColor="BDD7EE")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TINT     = PatternFill("solid", fgColor="F2F7FD")

IN_F   = Font(name=F, size=10, color="0000FF")      # 직접 입력
FX_F   = Font(name=F, size=10, color="000000")      # 같은 시트 자동
LNK_F  = Font(name=F, size=10, color="008000")      # 타 시트 참조
HDR_F  = Font(name=F, size=10, bold=True, color="FFFFFF")
SEC_F  = Font(name=F, size=11, bold=True, color="FFFFFF")
TIT_F  = Font(name=F, size=16, bold=True, color=NAVY)
NOTE_F = Font(name=F, size=9, italic=True, color="808080")
LBL_F  = Font(name=F, size=10, bold=True, color="000000")
TXT_F  = Font(name=F, size=10)

_t = Side(style="thin", color="BFBFBF")
BOX = Border(left=_t, right=_t, top=_t, bottom=_t)
THIN = _t
MED  = Side(style="medium", color="2E75B6")

MONEY = '#,##0;(#,##0);-'
NUM   = '#,##0;;-'
DEC   = '0.0;;-'
GB    = '#,##0.0;;-'
DATE_ = 'yyyy-mm-dd'
PCT   = '0.0%'
PCT0  = '0%'

# ── 기존 데이터 범위 (v3)
W1, W2   = 5, 1004      # 업무일지
D1, D2   = 5, 1004      # 일일실시사항
K1, K2   = 5, 64        # 사건마스터
P1, P2   = 5, 24        # 코드설정 담당자
CS1, CS2 = 5, 21        # 코드설정 세부유형
TM1, TM2 = 5, 14        # 코드설정 팀
RQ1, RQ2 = 5, 14        # 코드설정 요청처
ST1, ST2 = 5, 7         # 코드설정 진행상태
DG1, DG2 = 5, 14        # 코드설정 일일구분
JB1, JB2 = 5, 7         # 코드설정 직무구분
EQ_F, EQ_L = 6, 105     # 장비 대장
LN_F, LN_L = 109, 508   # 불출 이력

# ── v4 신규 범위
PJ1, PJ2 = 5, 34        # 사업마스터 (30건)
PL1, PL2 = 5, 404       # 사업추진일지 (400건)
SG1, SG2 = 5, 14        # 코드설정 사업구분 (10)
MD1, MD2 = 5, 18        # 코드설정 매체 종류 (14)
AS1, AS2 = 5, 11        # 코드설정 분석 상태 (7)
SP1, SP2 = 6, 305       # 지원 대장 (300건)
PS1, PS2 = 5, 10        # 코드설정 추진단계 (6)
JOBS = ["출장", "분석", "교육"]

# ── 업무일지 열 (v3 확정)
WC = {"연번":"A","시작일":"B","종료일":"C","직무구분":"D","세부유형":"E","사건번호":"F","회차":"G",
      "사건명":"H","요청처":"I","주담당":"J","투입인원":"Z","명단":"AA","소속팀":"AB","장소":"AC",
      "매체":"AD","데이터량":"AE","소요일수":"AF","진행상태":"AG","비고":"AH"}
# ── 일일실시사항 열
DC = {"연번":"A","일자":"B","요일":"C","담당자":"D","소속팀":"E","업무구분":"F","사건번호":"G",
      "내용":"H","결과":"I","진행상태":"J","특이사항":"K"}

def R(sheet, col, a, b):
    """시트의 절대 범위 문자열"""
    q = f"'{sheet}'" if " " in sheet or "·" in sheet else sheet
    return f"{q}!${col}${a}:${col}${b}"

# 자주 쓰는 범위
WB_ = R("업무일지", WC["시작일"], W1, W2)
WCc = R("업무일지", WC["종료일"], W1, W2)
WD_ = R("업무일지", WC["직무구분"], W1, W2)
WE_ = R("업무일지", WC["세부유형"], W1, W2)
WF_ = R("업무일지", WC["사건번호"], W1, W2)
WJ_ = R("업무일지", WC["주담당"], W1, W2)
WZ_ = R("업무일지", WC["투입인원"], W1, W2)
WAA = R("업무일지", WC["명단"], W1, W2)
WAB = R("업무일지", WC["소속팀"], W1, W2)
WAD = R("업무일지", WC["매체"], W1, W2)
WAF = R("업무일지", WC["소요일수"], W1, W2)
WAG = R("업무일지", WC["진행상태"], W1, W2)
DB_ = R("일일실시사항", DC["일자"], D1, D2)
DD_ = R("일일실시사항", DC["담당자"], D1, D2)
DF_ = R("일일실시사항", DC["업무구분"], D1, D2)
DJ_ = R("일일실시사항", DC["진행상태"], D1, D2)

def title(ws, text, note):
    ws["A1"] = text; ws["A1"].font = TIT_F
    ws.row_dimensions[1].height = 24
    ws["A2"] = note; ws["A2"].font = NOTE_F
    ws.row_dimensions[2].height = 26

def hdr(ws, row, labels, start=1, height=30):
    for j, h in enumerate(labels, start=start):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill, c.border = HDR_F, HDR_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height

def section(ws, row, text, span=10):
    c = ws.cell(row=row, column=1, value=text)
    c.font = SEC_F
    for j in range(1, span + 1):
        ws.cell(row=row, column=j).fill = SEC_FILL
        ws.cell(row=row, column=j).border = BOX
    ws.row_dimensions[row].height = 20

def put(ws, cell, val, font=FX_F, fmt=None, fill=None, align=None, border=True):
    c = ws[cell]; c.value = val; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if border: c.border = BOX
    c.alignment = Alignment(horizontal=align or "center", vertical="center",
                            wrap_text=(align == "wrap"))
    return c

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def fresh_sheet(wb, name):
    """같은 이름의 시트가 있으면 지우고 새로 만든다 (스테이지 재실행 대비)"""
    if name in wb.sheetnames:
        del wb[name]
    return wb.create_sheet(name)

def dv_name(ws, rng, name):
    d = DataValidation(type="list", formula1=name, allow_blank=True, showDropDown=False)
    ws.add_data_validation(d); d.add(rng)

def dv_inline(ws, rng, items):
    d = DataValidation(type="list", formula1=f'"{items}"', allow_blank=True, showDropDown=False)
    ws.add_data_validation(d); d.add(rng)
