# -*- coding: utf-8 -*-
"""v4 정적 검증 — 수식 문법 · 한셀 호환 · 드롭다운 · 참조 정합성"""
import sys, re, collections
import openpyxl

path = sys.argv[1] if len(sys.argv) > 1 else "부서_업무일지_v4.xlsx"
wb = openpyxl.load_workbook(path)
sheets = set(wb.sheetnames)
names  = {n.name for n in wb.defined_names.values()}
bad = collections.defaultdict(list)
n_f = 0

SHEETREF = re.compile(r"(?:'([^']+)'|([가-힣][가-힣\w]*))!")
CELLREF  = re.compile(r"(?<![A-Z0-9_.!$])\$?([A-Z]{1,3})\$?(\d+)(?![\d(])")
FN       = re.compile(r"(?<![A-Z._$])([A-Z]{2,12})\s*\(")
BANNED   = {"XLOOKUP","XMATCH","UNIQUE","FILTER","SEQUENCE","SORT","LET","LAMBDA",
            "TEXTJOIN","MINIFS","MAXIFS","IFS","CONCAT","SWITCH","AGGREGATE","TEXTSPLIT"}
VER = {"IF":"5","SUM":"5","MAX":"5","MIN":"5","COUNT":"5","INDEX":"5","MATCH":"5","AND":"5",
       "OR":"5","NOT":"5","MID":"5","TRIM":"5","SUBSTITUTE":"5","REPT":"5","TEXT":"5",
       "DATE":"5","TODAY":"5","WEEKDAY":"5","CHOOSE":"5","MONTH":"5","DAY":"5","LEN":"5",
       "LEFT":"5","RIGHT":"5","ABS":"5","ROUND":"5","LARGE":"5","SMALL":"5","SEARCH":"5",
       "ISNUMBER":"5","SUMPRODUCT":"97","COUNTIF":"97","COUNTBLANK":"97","COUNTA":"97",
       "COUNTIFS":"2007","SUMIFS":"2007","IFERROR":"2007","EOMONTH":"2007"}
used = collections.Counter()

def args(s):
    out, depth, cur, q = [], 0, "", False
    for ch in s:
        if ch == '"': q = not q
        if not q:
            if ch == "(": depth += 1
            elif ch == ")":
                if depth == 0: break
                depth -= 1
            elif ch == "," and depth == 0:
                out.append(cur); cur = ""; continue
        cur += ch
    out.append(cur); return out

RNG = re.compile(r"^(?:'[^']+'|[가-힣\w]+)?!?\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)$")
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str) or not v.startswith("="): continue
            n_f += 1; at = f"{ws.title}!{c.coordinate}"
            for m in SHEETREF.finditer(v):
                nm = m.group(1) or m.group(2)
                if nm and nm not in sheets:
                    bad["없는 시트 참조"].append(f"{at}: '{nm}'")
            for m in FN.finditer(v):
                fn = m.group(1); used[fn] += 1
                if fn in BANNED: bad["한셀 비호환 함수"].append(f"{at}: {fn}")
                if "_xlfn" in v: bad["_xlfn 접두사"].append(at)
            loc = re.sub(r"'[^']+'!|[가-힣][가-힣\w]*!", "!", v)   # 치환 문자도 제외 대상이어야 함
            for m in CELLREF.finditer(loc):
                if f"{m.group(1)}{m.group(2)}" == c.coordinate:
                    bad["자기참조"].append(f"{at}: {v[:70]}"); break
            for fn, sum_first in (("COUNTIFS", False), ("SUMIFS", True)):
                i = 0; up = v.upper()
                while True:
                    k = up.find(fn + "(", i)
                    if k < 0: break
                    aa = args(v[k + len(fn) + 1:])
                    hs = set()
                    for idx, a in enumerate(aa):
                        if sum_first and idx == 0: pass
                        elif (idx - (1 if sum_first else 0)) % 2 != 0: continue
                        mm = RNG.match(a.strip())
                        if mm: hs.add(int(mm.group(4)) - int(mm.group(2)) + 1)
                    if len(hs) > 1: bad[f"{fn} 범위 높이 불일치"].append(f"{at}: {sorted(hs)}")
                    i = k + 1

for ws in wb.worksheets:
    for d in ws.data_validations.dataValidation:
        f1 = (d.formula1 or "").strip()
        if f1.startswith('"') or d.type != "list": continue
        if f1 not in names:
            bad["드롭다운 원본 미해결"].append(f"{ws.title} {str(d.sqref)[:18]} ← {f1}")

print(f"파일: {path}")
print(f"시트 {len(wb.worksheets)}개 · 수식 {n_f:,}개 · 이름정의 {len(names)}개")
print()
late = sorted({f for f in used if VER.get(f, "?") not in ("5", "97", "2007")})
print("── 함수 버전 ──")
print(f"  사용 {len(used)}종 / 2007년 이후 도입: ", end="")
print(sorted({f for f in used if VER.get(f) == "2007"}))
print(f"  버전 미확인·위험: {late or '없음 ✅'}")
print()
if bad:
    print("── 문제 ──")
    for k, v in bad.items():
        print(f"  ✖ {k} ({len(v)}건)")
        for x in v[:4]: print(f"      {x}")
        if len(v) > 4: print(f"      … 외 {len(v)-4}건")
else:
    print("✅ 정적 검사 통과 — 문제 없음")
